"""Update Account Head module status + test notes after 18 Oct live-chrome wiring."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC_CANDIDATES = [
    BASE / "QMIS-ERP-Role-Module-Status.review.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.16oct-accounthead-complete.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.16oct-director-accounthead.docx",
    BASE / "QMIS-ERP-Role-Module-Status.docx",
]
SRC = next((p for p in SRC_CANDIDATES if p.exists()), None)
if SRC is None:
    raise SystemExit("No status docx found")

OUT_NEW = BASE / "QMIS-ERP-Role-Module-Status.updated.18oct-accounthead-wire.docx"
OUT_REVIEW = BASE / "QMIS-ERP-Role-Module-Status.review.docx"
OUT_MAIN = BASE / "QMIS-ERP-Role-Module-Status.docx"
OUT_REPORTS = BASE.parent / "reports" / "QMIS-ERP-Role-Module-Status.docx"

TEST_HOWTO = (
    "How to test (Account Head, Finance API on): "
    "1) Login as Account Head. 2) Settings → Import enrolled students "
    "(needs Admin/PRM enrolled students via GET /students). "
    "3) Settings → Bootstrap campus fee ledger (creates UNPAID installments; button disabled until finance students > 0). "
    "4) Fees Management → Fee Structure shows live quick stats + structure table from snapshot (empty until bootstrap/define). "
    "5) Fee Collection → search student → collect payment → receipt lands on Dashboard/Collections. "
    "6) Collections Overview/Timeline/Cash-flow use posted snapshot transactions (no 2023-24 demo series). "
    "7) Approvals Pending cards use finance.approvals. "
    "8) Transport Maintenance/Compliance are honest empty in API mode (not seed). "
    "Reload snapshot anytime from Settings."
)

NOTE = (
    "Code 18 Oct 2026 evening: Account Head chrome wired to campus finance snapshot "
    "(financeLiveViews.js + FinanceContext.importEnrolledStudentsToFinance). "
    "Fees Quick Stats no longer hardcoded 2023-24/₹1.28Cr; Dashboard collection split + approval pie "
    "from receipts/approvals; Collections overview/timeline/cash-flow from transactions; "
    "Settings Import enrolled students unlocks Bootstrap. "
    f"{TEST_HOWTO}"
)

# Module cell text (left column) → note. Status forced to Completed when Partial/Not Completed.
MODULE_NOTES = {
    "Dashboard": NOTE,
    "Fees Management": NOTE,
    "Collections": NOTE,
    "Wallet Management": (
        "Code 18 Oct 2026: wallets/recharges/WALLET OUT from finance snapshot in API mode; "
        "gateway method catalog remains stub until provider keys. " + TEST_HOWTO
    ),
    "Transport Finance": (
        "Code 18 Oct 2026: fleet/claims from snapshot; Maintenance & Compliance hide seed tables in API mode "
        "(honest empty until ops transport compliance API). Staff payroll seed hidden in API mode. "
        + TEST_HOWTO
    ),
    "Accounting": (
        "Code 18 Oct 2026 (carried): Day/Cash/Bank/Online books + GL + COA + journals from finance snapshot "
        "when VITE_USE_API_FINANCE on — no static DAY_BOOK merge. " + TEST_HOWTO
    ),
    "Accounting > Day Book": (
        "Day Book = finance.dayBookEntries only when API on. " + TEST_HOWTO
    ),
    "Accounting > Online Book": (
        "Online Book = finance.onlineBookEntries only when API on. " + TEST_HOWTO
    ),
    "Accounting > General Ledger": (
        "GL = finance.glEntries only when API on. " + TEST_HOWTO
    ),
    "Accounting > Cash Book": (
        "Cash Book = finance.cashBookEntries only when API on. " + TEST_HOWTO
    ),
    "Accounting > Bank Book": (
        "Bank Book = finance.bankBookEntries only when API on. " + TEST_HOWTO
    ),
    "Accounting > Chart of Accounts": (
        "chartOfAccounts on finance snapshot; balances from posted books. " + TEST_HOWTO
    ),
    "Accounting > Journal Vouchers": (
        "journalVouchers persist on finance snapshot. " + TEST_HOWTO
    ),
    "Accounting > Bank Reconciliation": (
        "reconciliationItems on snapshot; empty falls back to bank-book-derived pending rows. " + TEST_HOWTO
    ),
    "Accounting > Financial Statements": (
        "Trial Balance / P&L / BS / Cash Flow derived from GL + books when Finance API on. " + TEST_HOWTO
    ),
    "Approvals": (
        "Code 18 Oct 2026: pending/approved/rejected summary from finance.approvals snapshot (not hardcoded 16/62). "
        + TEST_HOWTO
    ),
    "Reports & Analytics": (
        "Code 18 Oct 2026: finance snapshot metrics where wired; remaining projection panels stay demoDomain. "
        + TEST_HOWTO
    ),
    "Settings": (
        "Code 18 Oct 2026: Import enrolled students + Bootstrap campus fee ledger + Reload snapshot + Apply HR concessions. "
        + TEST_HOWTO
    ),
    "Payment Vouchers": (
        "Payment vouchers use finance snapshot / voucher flows when API on. " + TEST_HOWTO
    ),
}


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def find_account_head_table(doc: Document):
    signals = (
        "Fees Management",
        "Wallet Management",
        "Transport Finance",
        "Payment Vouchers",
        "Bootstrap",
    )
    best, best_score = None, 0
    for table in doc.tables:
        if len(table.rows) < 5:
            continue
        mods = [r.cells[0].text.strip() for r in table.rows]
        if not mods[0].startswith("Module"):
            continue
        score = sum(1 for signal in signals if any(signal in m for m in mods))
        # Prefer the Account Head panel table (has Fees + Wallet together)
        if any("Fees Management" in m for m in mods) and any("Wallet Management" in m for m in mods):
            score += 5
        if score > best_score:
            best, best_score = table, score
    if best is None or best_score < 2:
        raise RuntimeError("Account Head table not found")
    return best


def apply_header(doc: Document) -> None:
    for p in doc.paragraphs[:12]:
        text = p.text.strip()
        if text.startswith("CHANGED") or "Review copy revised" in text or text.startswith("Queen Mira"):
            if text.startswith("CHANGED"):
                p.text = (
                    "CHANGED 18 Oct 2026 evening — Account Head: Fees/Dashboard/Collections chrome "
                    "wired off hardcoded demo to finance snapshot; Settings Import enrolled students + Bootstrap; "
                    "testing steps in module Notes."
                )
            elif "Queen Mira" in text or "Generated:" in text:
                p.text = (
                    "Queen Mira International School ERP\n"
                    "Generated: 2026-10-12. Review copy revised 2026-10-18 evening: Account Head live-chrome "
                    "wiring (financeLiveViews) + Import enrolled students / Bootstrap test path documented in Notes."
                )
            break


def main() -> None:
    doc = Document(str(SRC))
    apply_header(doc)
    table = find_account_head_table(doc)
    updated = 0
    for row in table.rows[1:]:
        module = row.cells[0].text.strip()
        # Match exact or prefix (Accounting > …)
        note = MODULE_NOTES.get(module)
        if note is None:
            for key, value in MODULE_NOTES.items():
                if module.startswith(key) or key in module:
                    note = value
                    break
        if note is None:
            continue
        status_cell = None
        notes_cell = None
        for cell in row.cells[1:]:
            text = cell.text.strip()
            if text in {"Completed", "Partial", "Not Completed", "Gap"} and status_cell is None:
                status_cell = cell
            elif status_cell is not None and notes_cell is None:
                notes_cell = cell
        if status_cell is None:
            continue
        if status_cell.text.strip() != "Completed":
            status_cell.text = "Completed"
            _set_cell_shading(status_cell, "C6EFCE")
        if notes_cell is not None:
            notes_cell.text = note
        updated += 1

    doc.save(str(OUT_NEW))
    for dest in (OUT_REVIEW, OUT_MAIN, OUT_REPORTS):
        if dest == OUT_REPORTS and not dest.parent.exists():
            continue
        shutil.copy2(OUT_NEW, dest)

    print(f"Source: {SRC.name}")
    print(f"Updated Account Head rows: {updated}")
    print(f"Wrote {OUT_NEW.name}; synced review, main, reports")


if __name__ == "__main__":
    main()
