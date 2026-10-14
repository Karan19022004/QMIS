"""Mark PRM (Front Office) modules Completed after API wiring + await create/update fixes."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).resolve().parent

# Prefer the latest PRM sheet; fall back to evening then base.
for candidate in (
    BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-prm.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-evening.docx",
    BASE / "QMIS-ERP-Role-Module-Status.docx",
):
    if candidate.exists():
        SRC = candidate
        break
else:
    raise FileNotFoundError("No role-module status Word source found")

OUTS = [
    BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-prm.docx",
    BASE / "QMIS-ERP-Role-Module-Status.docx",
    BASE / "QMIS-ERP-Role-Module-Status.review.14oct.docx",
    BASE / "QMIS-ERP-Role-Module-Status.ADMIN-UPDATED-14oct.docx",
]

UPDATES = {
    "User Database > Teacher Database": (
        "Completed",
        "PRM Teacher Database uses EmployeeDirectory API (GET /hr/portal/employees + /admin/users) with department filter Teacher; "
        "PRM granted hr.employees.read (14 Oct 2026). Read-only list.",
    ),
    "Student Transfer": (
        "Completed",
        "List + Add persist via /api/v1/front-office/student-transfers (0011, frontoffice.*). "
        "AddStudentTransfer awaits POST; verified live STR-* (14 Oct 2026).",
    ),
    "Student Re-Enrollment": (
        "Completed",
        "Persisted on /api/v1/front-office/student-re-enrollments (0011). Add/list use Front Office API when signed in.",
    ),
    "TC Request": (
        "Completed",
        "PRM list/view hydrate GET /api/v1/tc-requests; forward POST …/forward; reject POST …/reject-prm (tc.write). "
        "Same tc_requests table as Admin approval.",
    ),
    "Gate Pass > Student Gate Pass": (
        "Completed",
        "Persisted on /api/v1/front-office/student-gate-passes (0011). Add awaits API create; UI layout preserved.",
    ),
    "Gate Pass > Hostel Gate Pass": (
        "Completed",
        "Add/Edit await create/update on /api/v1/front-office/hostel-gate-passes (0011). Live HGP-* verified.",
    ),
    "Gate Pass > Material Gate Pass": (
        "Completed",
        "Add/Edit await create/update on /api/v1/front-office/material-gate-passes (0011). Live MGP-* verified.",
    ),
    "Gate Pass > Goods Received Pass": (
        "Completed",
        "Add/Edit await create/update on /api/v1/front-office/goods-received-passes (0011). Live GRP-* verified.",
    ),
    "My Attendance": (
        "Completed",
        "PRM My Attendance reads GET /hr/attendance with hr.self (granted to prm). Falls back to demo rows only when API auth is off.",
    ),
    "Announcement": (
        "Completed",
        "Shared AnnouncementListView + announcements API (roleKey=prm). announcements.read for all roles.",
    ),
    "Communication": (
        "Completed",
        "Shared CommunicationPage + /communication APIs (roleKey=prm).",
    ),
    "Calendar": (
        "Completed",
        "Shared AcademicCalendarPage hydrates GET /calendar (roleKey=prm).",
    ),
    "Notifications": (
        "Completed",
        "PRM Notifications page loads GET /notifications and marks read (same pattern as Admin).",
    ),
}


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def find_prm_table(doc: Document):
    for table in doc.tables:
        if len(table.rows) < 5:
            continue
        mods = [r.cells[0].text.strip() for r in table.rows]
        if not mods[0].startswith("Module"):
            continue
        if "Student Transfer" in mods and "Admission Enquiry" in mods and "Gate Pass > Student Gate Pass" in mods:
            return table
    raise RuntimeError("PRM table not found")


def apply(doc: Document) -> tuple[dict, dict]:
    table = find_prm_table(doc)
    for row in table.rows[1:]:
        mod = row.cells[0].text.strip()
        if mod not in UPDATES:
            continue
        status, note = UPDATES[mod]
        row.cells[1].text = status
        row.cells[2].text = note
        set_cell_shading(row.cells[1], "C6EFCE")

    counts = {"Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in table.rows[1:]:
        status = row.cells[1].text.strip()
        if status.startswith("Completed"):
            counts["Completed"] += 1
        elif status.startswith("Partial"):
            counts["Partial"] += 1
        elif status.startswith("Gap"):
            counts["Gap"] += 1
        else:
            counts["Not Completed"] += 1

    summary = doc.tables[1]
    for row in summary.rows[1:]:
        if row.cells[0].text.strip().startswith("PRM"):
            row.cells[1].text = str(sum(counts.values()))
            row.cells[2].text = str(counts["Completed"])
            row.cells[3].text = str(counts["Partial"])
            row.cells[4].text = str(counts["Not Completed"])
            row.cells[5].text = str(counts["Gap"])
            break

    totals = {"Modules": 0, "Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in summary.rows[1:]:
        role = row.cells[0].text.strip()
        if not role or role == "TOTAL":
            continue
        try:
            totals["Modules"] += int(row.cells[1].text.strip() or 0)
            totals["Completed"] += int(row.cells[2].text.strip() or 0)
            totals["Partial"] += int(row.cells[3].text.strip() or 0)
            totals["Not Completed"] += int(row.cells[4].text.strip() or 0)
            totals["Gap"] += int(row.cells[5].text.strip() or 0)
        except ValueError:
            continue
    for row in summary.rows:
        if row.cells[0].text.strip() == "TOTAL":
            row.cells[1].text = str(totals["Modules"])
            row.cells[2].text = str(totals["Completed"])
            row.cells[3].text = str(totals["Partial"])
            row.cells[4].text = str(totals["Not Completed"])
            row.cells[5].text = str(totals["Gap"])
            break
    return counts, totals


def main() -> None:
    doc = Document(str(SRC))
    counts, totals = apply(doc)
    for out in OUTS:
        try:
            doc.save(str(out))
            print("Wrote", out.name)
        except PermissionError:
            alt = out.with_name(out.stem + ".unlocked.docx")
            doc.save(str(alt))
            print("LOCKED", out.name, "-> wrote", alt.name)
    print("Source", SRC.name)
    print("PRM counts", counts)
    print("TOTAL", totals)


if __name__ == "__main__":
    main()
