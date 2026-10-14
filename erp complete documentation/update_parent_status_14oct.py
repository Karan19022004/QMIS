"""Mark Parent modules Completed after verified API wiring (start with TC Request)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent

for candidate in (
    BASE / "QMIS-ERP-Role-Module-Status.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-prm.docx",
):
    if candidate.exists():
        SRC = candidate
        break
else:
    raise FileNotFoundError("No role-module status Word source found")

OUTS = [
    BASE / "QMIS-ERP-Role-Module-Status.docx",
    BASE / "QMIS-ERP-Role-Module-Status.review.14oct.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-parent.docx",
]

UPDATES = {
    "TC Request": (
        "Completed",
        "Parent/Student TC list/add/view hydrate /api/v1/tc-requests; create via POST (tc.write). "
        "Backend scopes parent list/create to linked children; parents cannot forward/reject-prm. "
        "Verified pytest test_parent_tc_requests (14 Oct 2026).",
    ),
}


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def find_parent_table(doc: Document):
    for table in doc.tables:
        if len(table.rows) < 10:
            continue
        mods = [r.cells[0].text.strip() for r in table.rows]
        if not mods[0].startswith("Module"):
            continue
        if (
            "TC Request" in mods
            and "Fees Payment > Academic Payment" in mods
            and "Transport > Track Bus" in mods
            and "Dashboard" in mods
        ):
            return table
    raise RuntimeError("Parent table not found")


def apply(doc: Document) -> tuple[dict, dict]:
    table = find_parent_table(doc)
    for row in table.rows[1:]:
        mod = row.cells[0].text.strip()
        if mod not in UPDATES:
            continue
        status, note = UPDATES[mod]
        row.cells[1].text = status
        row.cells[2].text = note
        set_cell_shading(row.cells[1], "C6EFCE")

    counts = {"Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in table.rows[1:]:
        status = row.cells[1].text.strip()
        if status.startswith("Completed"):
            counts["Completed"] += 1
        elif status.startswith("Partial"):
            counts["Partial"] += 1
        elif status.startswith("Gap"):
            counts["Gap"] += 1
        else:
            counts["Not Completed"] += 1

    summary = doc.tables[1]
    for row in summary.rows[1:]:
        if row.cells[0].text.strip().startswith("Parent"):
            row.cells[1].text = str(sum(counts.values()))
            row.cells[2].text = str(counts["Completed"])
            row.cells[3].text = str(counts["Partial"])
            row.cells[4].text = str(counts["Not Completed"])
            row.cells[5].text = str(counts["Gap"])
            break

    totals = {"Modules": 0, "Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in summary.rows[1:]:
        role = row.cells[0].text.strip()
        if not role or role == "TOTAL":
            continue
        try:
            totals["Modules"] += int(row.cells[1].text.strip() or 0)
            totals["Completed"] += int(row.cells[2].text.strip() or 0)
            totals["Partial"] += int(row.cells[3].text.strip() or 0)
            totals["Not Completed"] += int(row.cells[4].text.strip() or 0)
            totals["Gap"] += int(row.cells[5].text.strip() or 0)
        except ValueError:
            continue
    for row in summary.rows:
        if row.cells[0].text.strip() == "TOTAL":
            row.cells[1].text = str(totals["Modules"])
            row.cells[2].text = str(totals["Completed"])
            row.cells[3].text = str(totals["Partial"])
            row.cells[4].text = str(totals["Not Completed"])
            row.cells[5].text = str(totals["Gap"])
            break
    return counts, totals


def main() -> None:
    doc = Document(str(SRC))
    counts, totals = apply(doc)
    for out in OUTS:
        try:
            doc.save(str(out))
            print("Wrote", out.name)
        except PermissionError:
            alt = out.with_name(out.stem + ".unlocked.docx")
            doc.save(str(alt))
            print("LOCKED", out.name, "-> wrote", alt.name)
    print("Source", SRC.name)
    print("Parent counts", counts)
    print("TOTAL", totals)


if __name__ == "__main__":
    main()
