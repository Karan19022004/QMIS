"""Write a clearly labelled Admin status Word + HTML proof file."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-evening.docx"
OUT_DOCX = BASE / "QMIS-ERP-Role-Module-Status.ADMIN-UPDATED-14oct.docx"
OUT_HTML = BASE / "QMIS-ERP-Role-Module-Status.ADMIN-UPDATED-14oct.html"

OLD = {
    "Dashboard": "Not Completed",
    "Documents > Student Documents": "Not Completed",
    "Documents > Employee Documents": "Not Completed",
    "Announcement": "Not Completed",
    "Calendar": "Not Completed",
    "Communication": "Not Completed",
    "TC Request Approval": "Not Completed",
    "Activities > Cultural": "Not Completed",
    "Activities > Sports": "Not Completed",
    "Activities > Chess Result": "Not Completed",
    "Activities > Competitions": "Not Completed",
    "Notifications": "Not Completed",
    "Class Details": "Completed (pending merge)",
    "Leave Request > My Leave Requests": "Partial",
    "Leave Request > Received Leave Requests": "Partial",
    "User Database > Student Database": "Completed (pending merge)",
    "Attendance > Students": "Completed (pending merge)",
    "Attendance > Employees": "Completed (pending merge)",
    "Escalation Management": "Completed",
}


def insert_paragraph_after(paragraph, text, *, bold=False, size=11, color=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p


def find_admin_table(doc: Document):
    for table in doc.tables:
        if len(table.rows) < 5:
            continue
        mods = [r.cells[0].text.strip() for r in table.rows[:5]]
        if mods[0].startswith("Module") and "Dashboard" in mods[1] and any(
            "Admission" in m for m in mods
        ):
            return table
    raise RuntimeError("Administrator table not found")


def main() -> None:
    doc = Document(str(SRC))
    title = doc.paragraphs[0]
    lines = [
        (
            True,
            "CHANGED 14 Oct 2026 evening — Administrator panel only (main PR #3). "
            "If you still see Not Completed for Announcement/Documents/Calendar, you are looking at an OLD file.",
        ),
        (
            True,
            "Administrator counts: was 11 Completed / 3 Partial / 12 Not Completed  →  "
            "now 19 Completed / 7 Partial / 0 Not Completed.",
        ),
        (
            False,
            "Completed now: Dashboard, Documents (Student+Employee), Announcement, Calendar, Notifications, "
            "Class Details, Attendance, Leave, Escalation, Student/Employee DB (merged).",
        ),
        (
            False,
            "Partial (still): Activities Cultural/Sports/Chess/Competitions, TC Request Approval, "
            "Communication, Roles (read-only).",
        ),
        (
            False,
            'Open section "2. Executive summary" then the Administrator table. '
            "Other role panels (Principal, Teacher, etc.) were NOT changed this evening — they will look the same.",
        ),
    ]
    for bold, text in reversed(lines):
        insert_paragraph_after(
            title,
            text,
            bold=bold,
            size=12 if bold else 11,
            color=RGBColor(0xC0, 0x00, 0x00) if bold else RGBColor(0x1A, 0x1A, 0x1A),
        )

    doc.save(str(OUT_DOCX))
    print("Wrote", OUT_DOCX, OUT_DOCX.stat().st_size)

    admin = find_admin_table(doc)
    html_rows = []
    for row in admin.rows[1:]:
        mod = row.cells[0].text.strip()
        new_status = row.cells[1].text.strip()
        note = row.cells[2].text.strip()
        before = OLD.get(mod, "(not in tonight's change list)")
        changed = mod in OLD and before != new_status
        if changed and new_status.startswith("Completed"):
            bg = "#d1fae5"
        elif changed:
            bg = "#fef3c7"
        else:
            bg = "#ffffff"
        html_rows.append(
            "<tr style=\"background:{bg}\"><td>{mod}</td><td>{before}</td>"
            "<td><b>{new}</b></td><td>{note}</td></tr>".format(
                bg=bg,
                mod=mod,
                before=before,
                new=new_status,
                note=note[:240],
            )
        )

    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<title>Admin Role Module Status — 14 Oct evening</title>
<style>
body {{ font-family: Segoe UI, Arial, sans-serif; margin: 24px; }}
h1 {{ color: #b91c1c; }}
.banner {{ background:#fee2e2; border:2px solid #b91c1c; padding:12px 16px; margin-bottom:16px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
th, td {{ border: 1px solid #ccc; padding: 8px; vertical-align: top; }}
th {{ background: #111827; color: #fff; }}
.ok {{ color: #047857; font-weight: 700; }}
</style></head><body>
<div class="banner">
<strong>This page is the proof of change for 14 Oct 2026 evening.</strong><br>
Administrator on main (PR #3): <span class="ok">19 Completed</span>, 7 Partial, 0 Not Completed.<br>
Green rows = status changed to Completed. Yellow = Partial (wired; browser create/approve not run).<br>
Word file: <code>QMIS-ERP-Role-Module-Status.ADMIN-UPDATED-14oct.docx</code>
</div>
<h1>Administrator module status — before vs after</h1>
<table>
<tr><th>Module</th><th>Before (old review)</th><th>After (this evening)</th><th>Evidence note</th></tr>
{''.join(html_rows)}
</table>
</body></html>
"""
    OUT_HTML.write_text(html, encoding="utf-8")
    print("Wrote", OUT_HTML)

    check = Document(str(OUT_DOCX))
    print("Top paragraphs:")
    for i, p in enumerate(check.paragraphs[:8]):
        if p.text.strip():
            print(i, p.text[:180])


if __name__ == "__main__":
    main()
