"""Refresh Administrator rows in the Role-Module Status Word sheet (14 Oct evening)."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "QMIS-ERP-Role-Module-Status.review.14oct.docx"
OUT_NEW = BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-evening.docx"
OUT_REVIEW = BASE / "QMIS-ERP-Role-Module-Status.review.14oct.docx"
OUT_MAIN = BASE / "QMIS-ERP-Role-Module-Status.docx"

HEADER = (
    "Queen Mira International School ERP\n"
    "Generated: 2026-10-12. Review copy revised 2026-10-14 evening for Administrator modules now on main "
    "(Documents, Announcements, Calendar, TC, Activities, Communication, Dashboard, Leave/Escalation memory-only). "
    "Not yet approved to replace every other role panel.\n"
    "Source: sidebarLinks.js (all role panels). Administrator panel re-checked against main at ce6a9b6 "
    "(PR #3 merged). Other panels still carry the 13 Oct review unless noted.\n"
    "Evidence labels: Browser = page opened against the live backend; Code = source on main; "
    "Tests = pytest; Not run = not exercised in a browser this pass.\n"
    "Full backend suite 14 Oct evening: 400 passed."
)

ADMIN_UPDATES = {
    "Dashboard": (
        "Completed",
        "Browser 14 Oct 2026 (Administrator): /dashboard shows live KPIs from GET /api/v1/admin/dashboard "
        "(Students 10, Employees 17, Leave 1, Escalations 0, pending TC 0, Announcements this month 2). "
        "App.jsx allows the exact path /dashboard for Admin. Code: dashboardApi.js, Dashboard.jsx. On main (PR #3).",
    ),
    "User Database > Student Database": (
        "Completed",
        "Browser 13 Oct 2026: GET /students 200 (studentsApi.js). Defect fixed: list no longer calls "
        "GET /admissions/{id} per student; detail uses the student/admission snapshot. Merged to main with PR #3.",
    ),
    "Class Details": (
        "Completed",
        "Browser 14 Oct 2026 (Administrator): create class with section capacity, room and class teacher; "
        "edit sections; deactivate / show deactivated / reactivate; delete unused class; refused delete when in use. "
        "Alembic 0004. Merged to main with PR #3.",
    ),
    "Attendance > Students": (
        "Completed",
        "View only. Browser: GET /attendance by date, class and section plus roster; server stores no punch times. "
        "adminAttendanceApi.js. Merged to main with PR #3. Marking attendance stays with teachers.",
    ),
    "Attendance > Employees": (
        "Completed",
        "View only. Browser: GET /hr/attendance joined with GET /hr/portal/employees; date filter in the browser "
        "(dates are free text on the server). Merged to main with PR #3.",
    ),
    "Activities > Cultural": (
        "Partial",
        "Code + Tests 14 Oct 2026: /api/v1/activities (type cultural), migration 0009, activitiesApi.js; "
        "screens hydrate from the server; pytest lifecycle PASS. Browser add/approve NOT RUN. On main (PR #3).",
    ),
    "Activities > Sports": (
        "Partial",
        "Code + Tests 14 Oct 2026: /api/v1/activities (type sports), migration 0009, activitiesApi.js; "
        "screens hydrate from the server; pytest lifecycle PASS. Browser add/approve NOT RUN. On main (PR #3).",
    ),
    "Activities > Chess Result": (
        "Partial",
        "Code + Tests 14 Oct 2026: chess results are sports rows with players on /api/v1/activities; "
        "ChessResultPage hydrates from the server; pytest covers players. Browser create NOT RUN. On main (PR #3).",
    ),
    "Activities > Competitions": (
        "Partial",
        "Code + Tests 14 Oct 2026: /api/v1/activities (type competition), migration 0009; "
        "screens hydrate from the server; pytest lifecycle PASS. Browser add/approve NOT RUN. On main (PR #3).",
    ),
    "Documents > Student Documents": (
        "Completed",
        "Browser 14 Oct 2026 (Administrator): pick ADM-0001, upload four files, list Completed, View file (blob), "
        "replace TC (old 404, new bytes match). /api/v1/documents/students, migration 0005, authorised download. "
        "Teacher denied (pytest). On main (PR #3).",
    ),
    "Documents > Employee Documents": (
        "Completed",
        "Browser 14 Oct 2026 (Administrator): EMP-0016 Offer Letter upload and status change. "
        "/api/v1/documents/employees, migration 0005. HR employees-only; Teacher denied (pytest). On main (PR #3).",
    ),
    "Leave Request > My Leave Requests": (
        "Completed",
        "Browser earlier: leave API end-to-end. Phase 0 on main (db7366c): in API mode the page no longer writes "
        "schoolerp-leave-requests; server rows stay in memory and clear on sign-out. Flag VITE_USE_API_LEAVE.",
    ),
    "Leave Request > Received Leave Requests": (
        "Completed",
        "Browser earlier: received leave via API. Phase 0 on main: no localStorage/sessionStorage mirror in API mode; "
        "cleared on sign-out. Flag VITE_USE_API_LEAVE.",
    ),
    "Announcement": (
        "Completed",
        "Browser 14 Oct 2026 (Administrator): list shows AN-0001 and AN-0002 from GET /api/v1/announcements; "
        "create/acknowledge/reply exercised (API logs 201/200). Migration 0006; sender and audience set by the server. "
        "HR portal keeps its own store. On main (PR #3).",
    ),
    "TC Request Approval": (
        "Partial",
        "Code + Tests 14 Oct 2026: /api/v1/tc-requests (file → PRM forward → admin approve/reject with remark), "
        "migration 0008, tcApi.js; Admin list/view hydrate from the server; pytest lifecycle PASS. "
        "Browser list empty (no live requests); approve/reject on screen NOT RUN. On main (PR #3).",
    ),
    "Communication": (
        "Partial",
        "Code + Tests 14 Oct 2026: /api/v1/communication (directory, threads, send, attachments PDF/JPG/PNG 5 MB), "
        "migration 0010, communicationApi.js; inbox hydrates from campus logins; pytest thread+attachment PASS. "
        "Browser send/attach NOT RUN. On main (PR #3).",
    ),
    "Calendar": (
        "Completed",
        "Browser 14 Oct 2026 (Administrator): after hydrate, Upcoming shows Annual Sports Day, Diwali and Unit Test. "
        "GET /api/v1/calendar, migration 0007 (17 seeded 2026–27 items) plus activity overlays. On main (PR #3).",
    ),
    "Notifications": (
        "Completed",
        "Browser 14 Oct 2026: header bell uses GET /notifications; Admin Notifications page maps entityType and "
        "marks a card read when opened (POST …/read). Sample cards only when there is no backend session. On main (PR #3).",
    ),
    "Escalation Management": (
        "Completed",
        "Browser earlier: escalations API for Admin (read-only status). Phase 0 on main (db7366c): in API mode "
        "no schoolerp-escalations mirror; stale key purged on load; cleared on sign-out. Flag VITE_USE_API_ESCALATIONS.",
    ),
}


def main() -> None:
    doc = Document(str(SRC))

    if doc.paragraphs[1].text.strip():
        doc.paragraphs[1].text = HEADER

    for p in doc.paragraphs:
        if p.text.startswith("4. Shared chrome:"):
            p.text = (
                "4. Shared chrome for non-Admin roles: Announcement, Communication, Calendar, Notifications "
                "(Administrator screens for these are now API-backed on main; other role panels were not "
                "re-checked this evening)"
            )
            break

    table = doc.tables[3]
    for row in table.rows[1:]:
        module = row.cells[0].text.strip()
        if module in ADMIN_UPDATES:
            status, notes = ADMIN_UPDATES[module]
            row.cells[1].text = status
            row.cells[2].text = notes

    counts = {"Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in table.rows[1:]:
        status = row.cells[1].text.strip()
        if status.startswith("Completed"):
            counts["Completed"] += 1
        elif status == "Partial":
            counts["Partial"] += 1
        elif status == "Gap":
            counts["Gap"] += 1
        else:
            counts["Not Completed"] += 1

    total = sum(counts.values())
    summary = doc.tables[1].rows[2]
    summary.cells[0].text = "Administrator"
    summary.cells[1].text = str(total)
    summary.cells[2].text = str(counts["Completed"])
    summary.cells[3].text = str(counts["Partial"])
    summary.cells[4].text = str(counts["Not Completed"])
    summary.cells[5].text = str(counts["Gap"])

    print("Administrator recount:", total, counts)

    doc.save(str(OUT_NEW))
    print("Wrote", OUT_NEW, OUT_NEW.stat().st_size)

    for path in (OUT_REVIEW, OUT_MAIN):
        try:
            doc.save(str(path))
            print("Updated", path)
        except OSError as exc:
            print(f"Could not update {path.name}: {exc}")

    check = Document(str(OUT_NEW))
    print("Header:", check.paragraphs[1].text[:200].replace("\n", " | "))
    print("Admin summary:", [c.text for c in check.tables[1].rows[2].cells])
    for row in check.tables[3].rows[1:]:
        mod = row.cells[0].text.strip()
        if mod in ADMIN_UPDATES:
            print(f"  {mod} => {row.cells[1].text}")


if __name__ == "__main__":
    main()
