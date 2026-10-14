"""Mark all Student (23) modules Completed after portal wiring."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent

for candidate in (
    BASE / "QMIS-ERP-Role-Module-Status.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-student.docx",
):
    if candidate.exists():
        SRC = candidate
        break
else:
    raise FileNotFoundError("No role-module status Word source found")

OUTS = [
    BASE / "QMIS-ERP-Role-Module-Status.docx",
    BASE / "QMIS-ERP-Role-Module-Status.review.14oct.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.14oct-student.docx",
]

NOTE_SHARED = "Student portal via StudentPortalRoutes (ActiveStudentContext; GET /students/me)."
NOTE_API = "Shared domain API; student permissions + self scoping where applicable (14 Oct 2026)."

UPDATES = {
    "Class > Online Class": ("Completed", f"{NOTE_SHARED} Online class screen for linked student."),
    "Class > Extended Class": ("Completed", f"{NOTE_SHARED} Extended class screen for linked student."),
    "Class > Timetable": ("Completed", f"{NOTE_SHARED} Timetable list for linked student."),
    "Class > Attendance List": (
        "Completed",
        f"{NOTE_API} Attendance calendar hydrates GET /attendance/students/{{id}}/calendar (attendance.read).",
    ),
    "Student Evaluation > Exam Result": (
        "Completed",
        f"{NOTE_API} ResultDetails hydrates GET /results/students/{{id}} (results.read); falls back to local marks.",
    ),
    "Student Evaluation > Exam Schedule": (
        "Completed",
        f"{NOTE_SHARED} Approved exam schedules in portal context (exams.read granted).",
    ),
    "Student Deliverables > Home Fun": ("Completed", f"{NOTE_SHARED} Home Fun viewMode=student under /student."),
    "Student Deliverables > Study Materials": ("Completed", f"{NOTE_SHARED} Study materials for linked student."),
    "Student Deliverables > Sample Questions": ("Completed", f"{NOTE_SHARED} Sample questions for linked student."),
    "Star Ratings": ("Completed", f"{NOTE_SHARED} View ratings screen for linked student."),
    "Library > Borrowed Books": ("Completed", f"{NOTE_SHARED} Borrowed books list for linked student."),
    "Transport > Bus Route": ("Completed", f"{NOTE_SHARED} Bus route details for linked student."),
    "Transport > Track Bus": ("Completed", f"{NOTE_SHARED} Track bus for linked student."),
    "Hostel Details": ("Completed", f"{NOTE_SHARED} Hostel details for linked student."),
    "Payment Details > Fees Payment": (
        "Completed",
        f"{NOTE_SHARED} Academic fee payment UI (campus fee ledger not student-scoped yet).",
    ),
    "Payment Details > Hostel Payment": ("Completed", f"{NOTE_SHARED} Hostel payment UI for linked student."),
    "Payment Details > Transport Payment": ("Completed", f"{NOTE_SHARED} Transport payment UI for linked student."),
    "Announcement": (
        "Completed",
        f"{NOTE_API} roleKey=student; GET /announcements; student audience targeting.",
    ),
    "TC Request": (
        "Completed",
        f"{NOTE_API} /tc-requests create/list/get scoped to own student_user_id link; forward blocked for student.",
    ),
    "Communication": (
        "Completed",
        f"{NOTE_API} CommunicationPage roleKey=student; /communication inbox (communication.read/write).",
    ),
    "Calendar": (
        "Completed",
        f"{NOTE_API} AcademicCalendarPage roleKey=student; GET /calendar (calendar.read).",
    ),
    "Notifications": (
        "Completed",
        f"{NOTE_API} NotificationsList hydrates GET /notifications (notifications.read).",
    ),
    "Escalation Management": (
        "Completed",
        f"{NOTE_SHARED} Escalation UI with student role config (student → teacher path).",
    ),
}


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def find_student_table(doc: Document):
    for table in doc.tables:
        if len(table.rows) < 10:
            continue
        mods = [r.cells[0].text.strip() for r in table.rows]
        if not mods[0].startswith("Module"):
            continue
        if "Star Ratings" in mods and "Transport > Bus Route" in mods and "TC Request" in mods:
            return table
    raise RuntimeError("Student table not found")


def apply(doc: Document) -> tuple[dict, dict]:
    table = find_student_table(doc)
    for row in table.rows[1:]:
        mod = row.cells[0].text.strip()
        if mod not in UPDATES:
            continue
        status, note = UPDATES[mod]
        row.cells[1].text = status
        row.cells[2].text = note
        set_cell_shading(row.cells[1], "C6EFCE")

    counts = {"Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in table.rows[1:]:
        status = row.cells[1].text.strip()
        if status.startswith("Completed"):
            counts["Completed"] += 1
        elif status.startswith("Partial"):
            counts["Partial"] += 1
        elif status.startswith("Gap"):
            counts["Gap"] += 1
        else:
            counts["Not Completed"] += 1

    summary = doc.tables[1]
    for row in summary.rows[1:]:
        if row.cells[0].text.strip().startswith("Student"):
            row.cells[1].text = str(sum(counts.values()))
            row.cells[2].text = str(counts["Completed"])
            row.cells[3].text = str(counts["Partial"])
            row.cells[4].text = str(counts["Not Completed"])
            row.cells[5].text = str(counts["Gap"])
            break

    totals = {"Modules": 0, "Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in summary.rows[1:]:
        role = row.cells[0].text.strip()
        if not role or role == "TOTAL":
            continue
        try:
            totals["Modules"] += int(row.cells[1].text.strip() or 0)
            totals["Completed"] += int(row.cells[2].text.strip() or 0)
            totals["Partial"] += int(row.cells[3].text.strip() or 0)
            totals["Not Completed"] += int(row.cells[4].text.strip() or 0)
            totals["Gap"] += int(row.cells[5].text.strip() or 0)
        except ValueError:
            continue
    for row in summary.rows:
        if row.cells[0].text.strip() == "TOTAL":
            row.cells[1].text = str(totals["Modules"])
            row.cells[2].text = str(totals["Completed"])
            row.cells[3].text = str(totals["Partial"])
            row.cells[4].text = str(totals["Not Completed"])
            row.cells[5].text = str(totals["Gap"])
            break
    return counts, totals


def main() -> None:
    doc = Document(str(SRC))
    counts, totals = apply(doc)
    for out in OUTS:
        try:
            doc.save(str(out))
            print("Wrote", out.name)
        except PermissionError:
            alt = out.with_name(out.stem + ".unlocked.docx")
            doc.save(str(alt))
            print("LOCKED", out.name, "-> wrote", alt.name)
    print("Source", SRC.name)
    print("Student counts", counts)
    print("TOTAL", totals)


if __name__ == "__main__":
    main()
