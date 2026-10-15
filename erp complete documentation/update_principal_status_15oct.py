"""Update Principal (21) module statuses after API + portal build (15 Oct 2026)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent

for candidate in (
    BASE / "QMIS-ERP-Role-Module-Status.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.15oct-principal.docx",
    BASE / "QMIS-ERP-Role-Module-Status.ADMIN-UPDATED-14oct.docx",
):
    if candidate.exists():
        SRC = candidate
        break
else:
    raise FileNotFoundError("No role-module status Word source found")

OUTS = [
    BASE / "QMIS-ERP-Role-Module-Status.docx",
    BASE / "QMIS-ERP-Role-Module-Status.review.14oct.docx",
    BASE / "QMIS-ERP-Role-Module-Status.updated.15oct-principal.docx",
]

NOTE_API = "Shared domain API with principal permissions (15 Oct 2026 build)."
NOTE_PORTAL = "Principal portal UI; shared Teacher/Common SoT (local until dedicated API)."

UPDATES = {
    "Task Management > Assign Tasks": (
        "Completed",
        f"{NOTE_API} Tasks API via TaskManagementRoutes.",
    ),
    "Task Management > My Tasks": (
        "Completed",
        f"{NOTE_API} Tasks API via TaskManagementRoutes.",
    ),
    "Star Ratings > Star of the Month": (
        "Partial",
        f"{NOTE_PORTAL} Staff SOM UI present; no ratings REST API yet.",
    ),
    "Star Ratings > Star of the Year": (
        "Partial",
        f"{NOTE_PORTAL} Staff SOY UI present; no ratings REST API yet.",
    ),
    "Examination Timetable": (
        "Completed",
        f"{NOTE_API} List hydrates GET /exams; create POSTs /exams + submit for Director approval.",
    ),
    "Class Timetable": (
        "Partial",
        f"{NOTE_PORTAL} Class timetable + change-request UI; no class-timetable REST module yet.",
    ),
    "User Database > Student Database": (
        "Completed",
        f"{NOTE_API} StudentDatabaseListView hydrates GET /students; search/status/DOB filters wired.",
    ),
    "User Database > Employee Database": (
        "Completed",
        f"{NOTE_API} EmployeeDatabaseListView campus directory GET /hr/portal/employees.",
    ),
    "Teacher Allocation": (
        "Partial",
        f"{NOTE_PORTAL} Allocation screens + localStorage SoT; no teacher-allocation API yet.",
    ),
    "Student Allocation": (
        "Completed",
        f"{NOTE_PORTAL} Allocator approval UI (StudentAllocationList mode=allocator).",
    ),
    "Extended Class": (
        "Completed",
        f"{NOTE_PORTAL} Reuses Teacher ExtraClass with principal list/add/view paths.",
    ),
    "Home Fun > Student Home Fun": (
        "Completed",
        f"{NOTE_PORTAL} Wraps Teacher HomeFun SoT at /principal/lms/student-lms (+ add/view).",
    ),
    "Home Fun > Teacher Home Fun": (
        "Completed",
        f"{NOTE_PORTAL} Wraps Teacher HomeFun SoT at /principal/lms/teacher-lms (+ add/view).",
    ),
    "Leave Request > My Leave Requests": (
        "Completed",
        f"{NOTE_API} MyLeaveRequestsList roleKey=principal.",
    ),
    "Leave Request > Received Leave Requests": (
        "Completed",
        f"{NOTE_API} ReceivedLeaveRequestsList roleKey=principal (leave.decide).",
    ),
    "My Attendance": (
        "Partial",
        f"{NOTE_PORTAL} Staff my-attendance UI + filters; no dedicated staff self-attendance API yet.",
    ),
    "Announcement": (
        "Completed",
        f"{NOTE_API} AnnouncementListView roleKey=principal (announcements.write).",
    ),
    "Communication": (
        "Completed",
        f"{NOTE_API} CommunicationPage roleKey=principal.",
    ),
    "Calendar": (
        "Completed",
        f"{NOTE_API} AcademicCalendarPage roleKey=principal; GET + Add Event POST (calendar.write).",
    ),
    "Notifications": (
        "Completed",
        f"{NOTE_API} Notifications hydrates GET /notifications.",
    ),
    "Escalation Management": (
        "Completed",
        f"{NOTE_API} EscalationListPage roleKey=principal.",
    ),
}


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def find_principal_table(doc: Document):
    for table in doc.tables:
        if len(table.rows) < 18:
            continue
        mods = [r.cells[0].text.strip() for r in table.rows]
        if not mods[0].startswith("Module"):
            continue
        if "Examination Timetable" in mods and "Teacher Allocation" in mods and "Student Allocation" in mods:
            return table
    raise RuntimeError("Principal table not found")


def apply(doc: Document) -> tuple[dict, dict]:
    table = find_principal_table(doc)
    for row in table.rows[1:]:
        mod = row.cells[0].text.strip()
        if mod not in UPDATES:
            continue
        status, note = UPDATES[mod]
        row.cells[1].text = status
        row.cells[2].text = note
        if status.startswith("Completed"):
            fill = "C6EFCE"
        elif status.startswith("Partial"):
            fill = "FFEB9C"
        elif status.startswith("Gap"):
            fill = "D9D9D9"
        else:
            fill = "FFC7CE"
        set_cell_shading(row.cells[1], fill)

    counts = {"Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in table.rows[1:]:
        status = row.cells[1].text.strip()
        if status.startswith("Completed"):
            counts["Completed"] += 1
        elif status.startswith("Partial"):
            counts["Partial"] += 1
        elif status.startswith("Gap"):
            counts["Gap"] += 1
        else:
            counts["Not Completed"] += 1

    summary = doc.tables[1]
    for row in summary.rows[1:]:
        if row.cells[0].text.strip().startswith("Principal"):
            row.cells[1].text = str(sum(counts.values()))
            row.cells[2].text = str(counts["Completed"])
            row.cells[3].text = str(counts["Partial"])
            row.cells[4].text = str(counts["Not Completed"])
            row.cells[5].text = str(counts["Gap"])
            break

    totals = {"Modules": 0, "Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for row in summary.rows[1:]:
        role = row.cells[0].text.strip()
        if not role or role == "TOTAL":
            continue
        try:
            totals["Modules"] += int(row.cells[1].text.strip() or 0)
            totals["Completed"] += int(row.cells[2].text.strip() or 0)
            totals["Partial"] += int(row.cells[3].text.strip() or 0)
            totals["Not Completed"] += int(row.cells[4].text.strip() or 0)
            totals["Gap"] += int(row.cells[5].text.strip() or 0)
        except ValueError:
            continue
    for row in summary.rows:
        if row.cells[0].text.strip() == "TOTAL":
            row.cells[1].text = str(totals["Modules"])
            row.cells[2].text = str(totals["Completed"])
            row.cells[3].text = str(totals["Partial"])
            row.cells[4].text = str(totals["Not Completed"])
            row.cells[5].text = str(totals["Gap"])
            break
    return counts, totals


def main() -> None:
    doc = Document(str(SRC))
    counts, totals = apply(doc)
    for out in OUTS:
        try:
            doc.save(str(out))
            print("Wrote", out.name)
        except PermissionError:
            alt = out.with_name(out.stem + ".unlocked.docx")
            doc.save(str(alt))
            print("LOCKED", out.name, "-> wrote", alt.name)
    print("Source", SRC.name)
    print("Principal counts", counts)
    print("TOTAL", totals)


if __name__ == "__main__":
    main()
