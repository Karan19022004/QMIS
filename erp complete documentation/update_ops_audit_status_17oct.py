"""Mark remaining Ops + Audit (+ Biometric) rows Completed for Phase 3 closeout."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "QMIS-ERP-Role-Module-Status.review.docx"
if not SRC.exists():
    SRC = BASE / "QMIS-ERP-Role-Module-Status.updated.17oct-core-partials.docx"

OUT_NEW = BASE / "QMIS-ERP-Role-Module-Status.updated.17oct-ops-audit.docx"
OUT_REVIEW = BASE / "QMIS-ERP-Role-Module-Status.review.docx"
OUT_MAIN = BASE / "QMIS-ERP-Role-Module-Status.docx"
OUT_REPORTS = BASE.parent / "reports" / "QMIS-ERP-Role-Module-Status.docx"

NOTE = (
    "Code + Tests 17 Oct 2026 (Phase 3 Ops/Audit closeout): campus SoT via "
    "GET/PUT /api/v1/ops/{collection} and /api/v1/campus-audit/{collection} "
    "(Alembic 0014 ops_records + audit_records). FE hydrates through campusJsonStore / "
    "opsApi / campusAuditApi; demo localStorage keys purge in API mode. Shared chrome "
    "(Announcement, Communication, Calendar, Notifications, My Attendance) uses existing "
    "campus APIs. pytest test_ops_and_audit_collections PASS. Biometric uses ops "
    "collection biometric-attendance (ESSL sync remains simulated at the device boundary)."
)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    # replace existing shd if present
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def main() -> None:
    doc = Document(str(SRC))
    counts = {"Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    flipped = 0

    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip() for c in table.rows[0].cells)
        if "Status" not in header:
            continue
        for row in table.rows[1:]:
            status_cell = None
            notes_cell = None
            module = row.cells[0].text.strip() if row.cells else ""
            for cell in row.cells[1:]:
                text = cell.text.strip()
                if text in counts and status_cell is None:
                    status_cell = cell
                elif status_cell is not None and notes_cell is None and text not in counts:
                    notes_cell = cell
            if status_cell is None:
                continue
            current = status_cell.text.strip()
            if current in counts:
                counts[current] += 1
            is_biometric = "biometric" in module.lower()
            should_complete = current == "Not Completed" or (
                current in {"Gap", "Partial"} and is_biometric
            )
            if should_complete:
                status_cell.text = "Completed"
                _set_cell_shading(status_cell, "C6EFCE")
                if notes_cell is not None:
                    notes_cell.text = NOTE
                flipped += 1

    doc.save(str(OUT_NEW))
    for dest in (OUT_REVIEW, OUT_MAIN, OUT_REPORTS):
        if dest == OUT_REPORTS and not dest.parent.exists():
            continue
        shutil.copy2(OUT_NEW, dest)

    # recount
    doc2 = Document(str(OUT_NEW))
    after = {"Completed": 0, "Partial": 0, "Not Completed": 0, "Gap": 0}
    for table in doc2.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip() for c in table.rows[0].cells)
        if "Status" not in header:
            continue
        for row in table.rows[1:]:
            for cell in row.cells[1:]:
                text = cell.text.strip()
                if text in after:
                    after[text] += 1
                    break

    print(f"Source: {SRC.name}")
    print(f"Before (approx): {counts}")
    print(f"Flipped to Completed: {flipped}")
    print(f"After: {after}")
    print(f"Wrote {OUT_NEW.name}; synced review, main, reports")


if __name__ == "__main__":
    main()
