# -*- coding: utf-8 -*-
"""Generate QMIS_ERP_Overall_Client_User_Manual.docx (+ PDF if possible). Run from repo root."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "QMIS_ERP_Overall_Client_User_Manual.docx"
PDF_PATH = OUT_DIR / "QMIS_ERP_Overall_Client_User_Manual.pdf"


def set_run(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run(run, bold=True, size={1: 16, 2: 14, 3: 12}.get(level, 11))
    return h


def p(doc, text, *, bold=False, size=11, space_after=8):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = para.add_run(text)
    set_run(run, bold=bold, size=size)
    return para


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(item)
        set_run(run, size=11)


def numbered(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Number")
        run = para.add_run(item)
        set_run(run, size=11)


def portal_block(
    doc,
    *,
    title,
    purpose,
    responsibilities,
    steps,
    info_needed,
    approval,
    outcome,
    mistakes,
):
    add_heading(doc, title, 2)
    add_heading(doc, "Portal Purpose", 3)
    p(doc, purpose)
    add_heading(doc, "Main Responsibilities", 3)
    bullets(doc, responsibilities)
    add_heading(doc, "Step-by-Step User Workflow", 3)
    numbered(doc, steps)
    add_heading(doc, "Information the User Needs", 3)
    bullets(doc, info_needed)
    add_heading(doc, "Approval and Handover", 3)
    p(doc, approval)
    add_heading(doc, "Expected Outcome", 3)
    p(doc, outcome)
    add_heading(doc, "Common Mistakes to Avoid", 3)
    bullets(doc, mistakes)


def workflow_block(doc, *, title, start, info, sequence, review, next_dept, check):
    add_heading(doc, title, 2)
    p(doc, f"Who starts: {start}", bold=False)
    p(doc, f"Information needed: {info}")
    add_heading(doc, "Sequence of actions", 3)
    numbered(doc, sequence)
    p(doc, f"Review / approval: {review}")
    p(doc, f"Next handover: {next_dept}")
    p(doc, f"What to check: {check}")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("QMIS ERP — Overall User Manual  |  Page ")
    set_run(run, size=9, color=RGBColor(0x55, 0x55, 0x55))
    # PAGE field
    fld = (
        '<?xml version="1.0"?>'
        '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr=" PAGE "/>'
    )
    from docx.oxml import parse_xml

    fp._p.append(parse_xml(fld))

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("QMIS ERP — Overall User Manual")
    set_run(r, bold=True, size=22, color=RGBColor(0x0C, 0x1E, 0x5B))
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("User Guide for School Management, Staff, Parents and Students")
    set_run(r, size=14, color=RGBColor(0x33, 0x33, 0x33))
    for line in (
        "Queen Mira International School",
        "Enterprise Resource Planning (ERP)",
        "",
        "Official client user guide",
        "For school leadership, academic and administrative staff,",
        "campus operations teams, parents and students",
    ):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(para.add_run(line), size=11)
    doc.add_page_break()

    # Purpose / how to use
    add_heading(doc, "1. Purpose of This Manual", 1)
    p(
        doc,
        "This manual explains how Queen Mira International School uses QMIS ERP for day-to-day "
        "school operations. It describes what each role does, how common school processes run from "
        "start to finish, and how departments hand work to one another.",
    )
    p(
        doc,
        "The guide is written for non-technical readers: school management, teaching and office "
        "staff, finance and HR teams, campus operations staff, parents and students.",
    )

    add_heading(doc, "2. How to Use This Manual", 1)
    bullets(
        doc,
        [
            "Start with the general sign-in instructions if you are new to the system.",
            "Use the overview of roles to find your portal.",
            "Open your portal chapter for step-by-step tasks.",
            "Use the end-to-end workflow chapters when a process involves more than one department.",
            "Follow the handover notes so you know who acts next.",
            "Use the final guidance section for safe daily use of accounts and school data.",
        ],
    )
    p(
        doc,
        "Contents (major sections): 1 Purpose · 2 How to use · 3 Sign-in · 4 Role overview · "
        "5 System administration · 6 Leadership and academics · 7 Admission and family portals · "
        "8 Finance · 9 HR · 10 Transport · 11 Campus operations · 12 Governance and audit · "
        "13 End-to-end workflows · 14 How departments connect · 15 General operating guidance.",
    )

    add_heading(doc, "3. General Sign-In and Account Use", 1)
    add_heading(doc, "3.1 Signing in", 2)
    numbered(
        doc,
        [
            "Open the school’s QMIS ERP sign-in page using the address provided by the school.",
            "Enter the email address assigned to your school account.",
            "Enter your password.",
            "Select Login.",
            "If the system asks you to change your password on first use, choose a strong new password "
            "and confirm it before continuing.",
            "After a successful sign-in, you are taken to the home area for your role.",
        ],
    )
    add_heading(doc, "3.2 Using your account safely", 2)
    bullets(
        doc,
        [
            "Use only the account issued to you. Do not share passwords.",
            "Sign out when you leave a shared computer.",
            "Check the academic year, class, section, student name and amounts before saving.",
            "If an action is not available on your screen, contact the responsible department—do not "
            "borrow another person’s login.",
            "Report incorrect student, fee or staff information to the department that owns that record.",
        ],
    )

    add_heading(doc, "4. Overview of User Roles", 1)
    p(
        doc,
        "QMIS ERP organises school work into role-based portals. Each person sees menus and tasks "
        "that match their responsibility. The table below summarises the main roles used by the school.",
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Role"
    hdr[1].text = "Primary school focus"
    rows = [
        ("Super Admin", "School-wide system and user oversight"),
        ("Admin", "School configuration, admissions support, academic masters, office coordination"),
        ("Director", "Leadership oversight, key academic approvals, monitoring"),
        ("Principal", "Academic administration, placement, examinations coordination"),
        ("Coordinator", "Grade-band academic coordination, placement support, marks review"),
        ("Teacher", "Assigned classes, attendance, lesson plans, marks entry"),
        ("PRM / Admission (Front Office)", "Enquiries, admission, enrolment, gate passes, office records"),
        ("Parent", "Linked children’s academics, fees, announcements, transport views"),
        ("Student", "Own timetable, attendance, learning items, published results"),
        ("Account Head", "Fee structures, collections, receipts, finance approvals"),
        ("Finance Assistant", "Routine fee review and permitted collections"),
        ("HR", "Employee records, onboarding, leave and staff administration"),
        ("Transport Manager", "Vehicles, routes, drivers, student transport assignments"),
        ("Transport Driver", "Assigned vehicle, duty and related operational updates"),
        ("Gate / Security", "Visitors, entry/exit and gate-pass handling"),
        ("Gate Keeper Manager", "Gate staffing and duty assignment oversight"),
        ("Librarian", "Catalogue, issue and return of books"),
        ("Stores / Stationery", "Stock receipts, issues and inventory monitoring"),
        ("Canteen", "Menu, orders and related canteen operations"),
        ("IT Support", "Support requests, progress and resolution records"),
        ("Housekeeping", "Schedules, duties and completion records"),
        ("Joint Director", "Cross-department oversight and assigned follow-up"),
        ("Joint Director Audit", "Audit assignments, findings and closure follow-up"),
        ("Process Auditor", "Process review, evidence and corrective actions"),
        ("Quality Auditor", "Quality checks, non-conformities and corrective actions"),
    ]
    for a, b in rows:
        cells = table.add_row().cells
        cells[0].text = a
        cells[1].text = b
    doc.add_paragraph()

    # --- Portals ---
    add_heading(doc, "5. System Administration Portals", 1)

    portal_block(
        doc,
        title="5.1 Super Admin",
        purpose="The Super Admin portal supports school-wide system administration and oversight. "
        "It is used for senior administrative control of users, important school-wide views and "
        "governance of access across the ERP.",
        responsibilities=[
            "Oversee system-level administration for the school.",
            "Manage authorised user and role administration within permitted boundaries.",
            "Review school-wide operational summaries and administrative modules assigned to this role.",
            "Support accountability for account activation and access-related administration.",
        ],
        steps=[
            "Sign in with your Super Admin account.",
            "Open the Super Admin dashboard to review the school-wide summary.",
            "Use User Creation (or equivalent user-administration menu) to create or review staff logins you are authorised to manage.",
            "Open the relevant oversight modules (for example finance overview, approvals, activity or audit-related views) as required by school policy.",
            "Confirm that account status and access changes are recorded before signing out.",
        ],
        info_needed=[
            "Full name, email and role of any staff account you are authorised to create or update.",
            "School policy on who may hold elevated access.",
            "Reason for activation, deactivation or password reset requests.",
        ],
        approval="Super Admin actions that create or change staff access should follow school authorisation. "
        "Day-to-day academic and finance work continues in the specialist portals; Super Admin does not replace "
        "those departments’ operational duties.",
        outcome="Authorised users can sign in with the correct role. Oversight screens show the expected school-wide information.",
        mistakes=[
            "Creating an account with the wrong role.",
            "Granting access without school authorisation.",
            "Changing another department’s operational records instead of asking the owning department.",
        ],
    )

    portal_block(
        doc,
        title="5.2 Admin",
        purpose="The Admin portal supports school configuration and day-to-day administrative coordination. "
        "Admin typically manages academic masters (classes and sections), supports admissions work, "
        "and coordinates office modules used across the school.",
        responsibilities=[
            "Maintain academic-year arrangements for classes and sections.",
            "Support admission lists and related Front Office records as assigned.",
            "Administer authorised staff or user records within Admin’s creatable set.",
            "Coordinate announcements, calendars, documents and office workflows available in the Admin menu.",
        ],
        steps=[
            "Sign in and open the Admin home area (commonly Admissions).",
            "Open Class Details to review or maintain classes and sections for the correct academic year.",
            "Open Admissions to review applications or continue enrolment tasks assigned to Admin.",
            "Use User Database or related menus for authorised student/staff directory work.",
            "Save changes and confirm the list or record shows the update after refresh.",
        ],
        info_needed=[
            "Current academic year.",
            "Class names, section names and capacity where required.",
            "Admission or student details provided by Front Office / PRM.",
        ],
        approval="Admin prepares and maintains school configuration used by academics, PRM and other departments. "
        "Specialist approvals (examinations, finance, HR leave) remain with those roles.",
        outcome="Classes, sections and administrative records are available for teaching, placement and admission teams.",
        mistakes=[
            "Editing the wrong academic year.",
            "Renaming or deactivating a class without checking dependent students.",
            "Entering admission data that belongs to PRM’s enquiry follow-up process without coordination.",
        ],
    )

    add_heading(doc, "6. Leadership and Academic Portals", 1)

    portal_block(
        doc,
        title="6.1 Director",
        purpose="The Director portal supports senior academic leadership: reviewing school activity, "
        "deciding key academic approvals (such as examination timetables and final marks publication) "
        "and monitoring assigned leadership tasks.",
        responsibilities=[
            "Review academic items awaiting Director decision.",
            "Approve or return examination timetable submissions according to school process.",
            "Provide final authorisation in the marks approval chain so results can be published.",
            "Monitor leadership tasks, broadcasts and related academic oversight menus.",
        ],
        steps=[
            "Sign in and open the Director home area.",
            "Open Approvals or the examination / marks review screens assigned to Director.",
            "Read the submission details (class, exam, marks batch, comments).",
            "Approve, or return with clear remarks if correction is required.",
            "Confirm the item leaves your pending list and that the next stakeholders can proceed.",
        ],
        info_needed=["Exam or marks submission details", "School academic calendar constraints", "Any remarks from Principal or Coordinator"],
        approval="Director is typically the final decision-maker for examination timetable approval and for publishing marks after Coordinator and Principal review.",
        outcome="Approved examinations and published results become available to authorised academic users and, once published, to permitted parent and student views.",
        mistakes=["Approving an incomplete marks batch", "Confusing draft marks with published results", "Omitting remarks when returning work for correction"],
    )

    portal_block(
        doc,
        title="6.2 Principal",
        purpose="The Principal portal supports academic administration for the campus: student placement oversight, "
        "teacher allocation, examination preparation, timetable coordination and mid-level academic approvals.",
        responsibilities=[
            "Oversee section placement and promotions as authorised.",
            "Coordinate teacher allocations and academic timetables.",
            "Create or submit official examination schedules for Director approval.",
            "Review marks batches at the Principal stage of the approval chain.",
            "Manage academic tasks, activities and staff-facing academic menus.",
        ],
        steps=[
            "Sign in and open Principal academic menus (tasks, allocation, examinations, placements as applicable).",
            "For placement: open the student allocation / placement area, select the enrolled student, choose section and confirm.",
            "For examinations: prepare the examination timetable and submit it for Director approval.",
            "For marks: open pending approval items, verify entries, then approve or return.",
            "Confirm saved placement or approval status after refresh.",
        ],
        info_needed=["Enrolled student list", "Section capacity", "Exam dates and subjects", "Marks batches awaiting review"],
        approval="Principal reviews work from Coordinators and Teachers and forwards examination and marks decisions toward Director where required.",
        outcome="Students are placed in sections; examination schedules move to Director; marks progress toward publication.",
        mistakes=["Placing a student in the wrong section", "Submitting an incomplete exam timetable", "Approving marks without checking subject totals"],
    )

    portal_block(
        doc,
        title="6.3 Coordinator",
        purpose="The Coordinator portal supports grade-band academic coordination: section assignment for enrolled "
        "students in scope, academic follow-up, and the first review step for many marks batches.",
        responsibilities=[
            "Assign enrolled students in your grade band to sections.",
            "Support promotion activities where authorised.",
            "Coordinate lesson-plan and academic activity follow-up in your scope.",
            "Enter or review marks and act as the first approver in the marks chain when required.",
            "Monitor attendance and examination follow-up for assigned grades.",
        ],
        steps=[
            "Sign in and open Section Assignment / Student Placement.",
            "Review students awaiting section assignment in your scope.",
            "Select a student, choose the section and save. Roll numbers are assigned by the school process.",
            "Open marks or approval menus to review submitted batches.",
            "Approve or return with remarks, then confirm the item status updates.",
        ],
        info_needed=["Grade band and academic year", "Section capacity", "Subject marks sheets", "Teacher submissions"],
        approval="Coordinator hands placed students to teaching teams and forwards marks to Principal after review.",
        outcome="Students appear in the correct section for teachers; marks batches advance in the approval chain.",
        mistakes=["Assigning outside your grade band", "Leaving students unplaced after enrolment", "Approving marks for the wrong examination"],
    )

    portal_block(
        doc,
        title="6.4 Teacher",
        purpose="The Teacher portal is the classroom workspace: assigned classes and subjects, timetable, lesson plans, "
        "attendance, assessment entry and student academic follow-up.",
        responsibilities=[
            "View assigned classes, subjects and timetable.",
            "Prepare and submit lesson plans for review.",
            "Mark class attendance.",
            "Enter assessment marks for authorised examinations and submit for approval.",
            "Revise returned marks batches when requested.",
            "Communicate academic updates through the menus available to Teachers.",
        ],
        steps=[
            "Sign in and open the Teacher dashboard.",
            "Select the class/subject for attendance or marks.",
            "Enter attendance or marks carefully against the correct student list.",
            "Save drafts where available, then submit for approval when complete.",
            "Monitor notifications for returns or approvals and revise if required.",
        ],
        info_needed=["Class and section", "Subject", "Examination name and maximum marks", "Student list for the session"],
        approval="Submitted lesson plans and marks move to Coordinator → Principal → Director as configured. Teachers do not publish results themselves.",
        outcome="Attendance and marks are stored for review; after full approval, published results can appear for students and parents.",
        mistakes=["Entering marks for the wrong exam or class", "Submitting incomplete sheets", "Editing after submission without following the revise process"],
    )

    add_heading(doc, "7. Admission and Family Portals", 1)

    portal_block(
        doc,
        title="7.1 PRM / Admission (Front Office)",
        purpose="PRM (Front Office) manages admission enquiries, applicant records, enrolment and related office "
        "services such as gate passes, transfers and re-enrolment coordination.",
        responsibilities=[
            "Register and follow up admission enquiries.",
            "Maintain admission lists and applicant information.",
            "Complete enrolment, including guardian details and parent/student account linkage where used.",
            "Hand enrolled students to academic placement.",
            "Manage gate-pass and Front Office movement records as assigned.",
        ],
        steps=[
            "Sign in to the Front Office / PRM portal.",
            "Open Admission Enquiry to capture or update an enquiry.",
            "Move the applicant to Admission List when ready and complete required fields.",
            "Use Enrol to create the student record and link parent/student accounts when required by school policy.",
            "Confirm the student appears as enrolled and inform academic staff for placement.",
            "For gate passes, open the relevant Gate Pass menu, enter details, save and share with Security as needed.",
        ],
        info_needed=["Applicant and guardian contact details", "Preferred class", "Identity and document details as per school policy", "Account emails for parent/student portals when linking logins"],
        approval="Enrolled students are handed to Principal/Coordinator for section placement. Gate passes may require further approval fields (for example hostel-related approvals) according to pass type.",
        outcome="Enquiry and admission records are complete; enrolled students are ready for placement; gate-pass records are available to Security.",
        mistakes=["Duplicate enquiries for the same child", "Enrolling without guardian contacts", "Handing over before enrolment is finished"],
    )

    portal_block(
        doc,
        title="7.2 Parent",
        purpose="The Parent portal gives guardians a secure view of linked children only: academics, attendance, "
        "timetable and published results, fees and receipts, announcements and transport information where provided.",
        responsibilities=[
            "Select the correct child when more than one child is linked.",
            "Review attendance, timetable and published examination results.",
            "Review fee status and complete permitted online payments where enabled.",
            "Read announcements and notifications.",
            "View transport information for the linked child where available.",
        ],
        steps=[
            "Sign in with the parent account issued after enrolment.",
            "Change the temporary password if prompted.",
            "Select the child on the child-selection screen if required.",
            "Open Dashboard, Attendance, Exam Result, Fees or Transport menus as needed.",
            "For fees, confirm the child and amount before paying.",
            "Sign out when finished.",
        ],
        info_needed=["Parent login email", "Child’s name and class", "Fee references for payments"],
        approval="Parents do not approve school academic records. Fee payments are confirmed through the school’s finance process.",
        outcome="Parents see only their linked children. Published results and fee receipts appear after school processes complete.",
        mistakes=["Selecting the wrong child", "Expecting unpublished draft marks to appear as results", "Sharing the parent login with others"],
    )

    portal_block(
        doc,
        title="7.3 Student",
        purpose="The Student portal shows each learner their own academic information: timetable, attendance, "
        "learning activities, examinations and published results, plus related personal academic history where available.",
        responsibilities=[
            "View personal timetable and attendance.",
            "Access learning activities and announcements meant for students.",
            "View examination schedules and published results for self only.",
            "Review fee or hostel-related personal information if exposed in the student menus.",
        ],
        steps=[
            "Sign in with the student account.",
            "Complete any required first-time password change.",
            "Open Class, Attendance, Examinations or Results menus.",
            "Confirm that the information shown matches your class and subjects.",
        ],
        info_needed=["Student login", "Current class and section"],
        approval="Students view published information; they do not approve marks or examinations.",
        outcome="The student sees only their own records. Published results appear after Director (final) academic approval.",
        mistakes=["Using another student’s device session", "Mistaking homework notes for official published results"],
    )

    add_heading(doc, "8. Finance Portals", 1)
    p(
        doc,
        "Finance work is split between Account Head and Finance Assistant so that fee configuration and "
        "high-level approvals remain with the Account Head, while routine collection can be performed by the Assistant.",
    )

    portal_block(
        doc,
        title="8.1 Account Head",
        purpose="The Account Head portal is the school’s finance desk for fee structures, collections, "
        "receipts, reconciliation, financial reports and finance approvals.",
        responsibilities=[
            "Define and maintain fee structures and related finance configuration.",
            "Oversee collections and outstanding balances.",
            "Issue and review receipts and financial records.",
            "Approve finance requests that require Head authority.",
            "Reconcile books and review reports.",
        ],
        steps=[
            "Sign in to the Account Head portal.",
            "Open Fees Management to review or update fee structures for the correct academic year.",
            "Open Collections or Fee Collection to record authorised payments.",
            "Open Approvals for claims or expense items awaiting Head decision.",
            "Review Reports and accounting books, then confirm figures after saving.",
        ],
        info_needed=["Academic year", "Fee heads and amounts", "Student identifiers", "Payment mode and reference"],
        approval="Account Head authorises structure changes and Head-level finance approvals. Routine collection may also be done by Finance Assistant.",
        outcome="Fee structures and receipts are stored in the campus finance records; approved items leave the pending approval list.",
        mistakes=["Changing fee structures mid-term without school authorisation", "Posting a payment to the wrong student", "Approving expenses without supporting details"],
    )

    portal_block(
        doc,
        title="8.2 Finance Assistant",
        purpose="The Finance Assistant portal supports operational fee work: reviewing ledgers, collecting permitted "
        "fees, issuing receipts and following up outstanding balances—without changing fee masters or giving Head approvals.",
        responsibilities=[
            "Review fee ledgers and outstanding balances.",
            "Collect permitted fees and generate receipts.",
            "Follow up routine collection activities.",
            "Hand restricted actions (fee structure changes, Head approvals, finance settings) to the Account Head.",
        ],
        steps=[
            "Sign in (you land in the shared finance workspace with Assistant permissions).",
            "Open Fees Management or Collections to locate the student and outstanding instalments.",
            "Record the permitted collection and confirm the receipt.",
            "If a fee structure change or approval is required, notify the Account Head instead of attempting restricted screens.",
            "Refresh the list to confirm the payment appears.",
        ],
        info_needed=["Student name / admission details", "Amount due", "Payment mode"],
        approval="Finance Assistant does not approve Head-level finance items and does not change fee structures. Those remain with Account Head.",
        outcome="Permitted collections and receipts are recorded. Restricted finance menus remain with Account Head.",
        mistakes=["Attempting to edit fee structures", "Collecting against the wrong instalment", "Treating a payment as complete before the receipt is saved"],
    )

    add_heading(doc, "9. Human Resources Portal", 1)
    portal_block(
        doc,
        title="9.1 HR",
        purpose="The HR portal manages employee information, onboarding, leave administration and related staff records.",
        responsibilities=[
            "Maintain employee profiles and staff directories information.",
            "Support onboarding and offer-related staff processes.",
            "Administer leave requests and staff attendance information where available.",
            "Produce staff administration reports used by leadership.",
        ],
        steps=[
            "Sign in to HR.",
            "Open Employees or Onboarding to create or update staff records.",
            "For leave, open Leave Management to review requests and record decisions within HR authority.",
            "Confirm the employee or leave status after saving.",
        ],
        info_needed=["Employee personal and contact details", "Designation and joining information", "Leave dates and type"],
        approval="Department heads may decide leave at their level; HR coordinates staff-wide leave and employee master data.",
        outcome="Staff records and leave decisions are available to authorised users.",
        mistakes=["Creating duplicate employee records", "Approving leave without checking balances", "Assigning the wrong designation"],
    )

    add_heading(doc, "10. Transport Portals", 1)
    portal_block(
        doc,
        title="10.1 Transport Manager",
        purpose="The Transport Manager portal runs fleet operations: vehicles, drivers, routes, stops, "
        "student transport assignments and duty monitoring.",
        responsibilities=[
            "Maintain vehicle and driver operational records.",
            "Define routes and stops.",
            "Assign students to transport routes.",
            "Plan driver duties and monitor transport operations.",
        ],
        steps=[
            "Sign in to Transport Manager.",
            "Update Vehicles, Drivers and Routes as needed.",
            "Open Student Transport to assign or change a student’s route.",
            "Assign duties and review operational requests (fuel/maintenance) according to school process.",
            "Confirm assignments after save.",
        ],
        info_needed=["Vehicle registration", "Driver details", "Route stops", "Student transport requirements"],
        approval="Transport Manager owns operational assignments. Drivers receive duty details. Finance may review transport expenses through finance processes.",
        outcome="Students and drivers have clear route and duty assignments.",
        mistakes=["Assigning a student to the wrong route", "Leaving a vehicle without a driver duty", "Updating the wrong vehicle record"],
    )

    portal_block(
        doc,
        title="10.2 Transport Driver",
        purpose="The Driver portal shows the driver’s assigned vehicle, route and duty information, and supports "
        "duty confirmation and reporting of vehicle or trip issues.",
        responsibilities=[
            "Review assigned vehicle and route.",
            "Confirm or follow daily duty instructions.",
            "Report vehicle health, fuel or maintenance concerns through the available request menus.",
        ],
        steps=[
            "Sign in to the Driver portal.",
            "Open My Duty and Vehicle details.",
            "Follow the assigned route and duty timing.",
            "Submit fuel or maintenance requests if required by school process.",
        ],
        info_needed=["Duty date", "Vehicle number", "Route name"],
        approval="Drivers report to Transport Manager. They do not reassign students or edit master routes.",
        outcome="Duty is visible and issues are logged for Transport Manager follow-up.",
        mistakes=["Starting duty on the wrong vehicle", "Submitting duplicate maintenance requests"],
    )
    p(
        doc,
        "Handover: Transport Manager creates routes, vehicles and student assignments; Drivers carry out the duty "
        "and report issues back to Transport Manager.",
    )

    add_heading(doc, "11. Campus Operations Portals", 1)

    portal_block(
        doc,
        title="11.1 Gate / Security",
        purpose="Gate and Security staff manage visitor registration, entry and exit records and gate-pass handling.",
        responsibilities=[
            "Register visitors and capture entry details.",
            "Check and process gate passes issued by Front Office or other authorised teams.",
            "Maintain relevant movement records for the campus gate.",
        ],
        steps=[
            "Sign in to the Gate Keeper portal (or Gate Keeper Manager portal if you assign duties).",
            "Open Visitors or Gate Pass menus.",
            "Verify identity and pass details before allowing entry or exit.",
            "Update status fields and save.",
        ],
        info_needed=["Visitor name and purpose", "Host/staff contact", "Gate-pass reference"],
        approval="Gate Keeper Manager may assign duties. Front Office prepares many pass records that Security executes.",
        outcome="Visitor and pass movements are recorded for campus safety.",
        mistakes=["Allowing entry without a valid pass", "Closing a pass against the wrong visitor"],
    )

    portal_block(
        doc,
        title="11.2 Librarian",
        purpose="The Librarian portal manages the book catalogue and the issue/return cycle for learners and staff borrowers.",
        responsibilities=["Register and catalogue books", "Issue books to borrowers", "Record returns", "Monitor borrowing records"],
        steps=[
            "Sign in to Librarian.",
            "Open Book Management to add or update catalogue entries.",
            "Use Issue/Return flows to lend and receive books.",
            "Confirm the borrower’s record updates after each transaction.",
        ],
        info_needed=["Book title and accession details", "Borrower identity", "Due dates as per library policy"],
        approval="Library operations are owned by the Librarian; academic staff may request titles through school channels.",
        outcome="Catalogue and borrower records stay current.",
        mistakes=["Issuing a book already on loan", "Returning against the wrong borrower"],
    )

    portal_block(
        doc,
        title="11.3 Stores / Stationery",
        purpose="Stores maintains item records, receives stock, issues materials and monitors inventory levels.",
        responsibilities=["Maintain item masters", "Record stock receipts", "Issue stock against requests", "Report discrepancies"],
        steps=[
            "Sign in to Stationery Store Manager.",
            "Update inventory items.",
            "Record receipts when stock arrives.",
            "Issue stock against approved requests.",
            "Review alerts for low stock or discrepancies.",
        ],
        info_needed=["Item name and unit", "Quantity", "Requesting department"],
        approval="Store requests may require Principal or Store Manager action depending on the request type.",
        outcome="Stock balances reflect receipts and issues.",
        mistakes=["Issuing more than available stock", "Posting receipt to the wrong item"],
    )

    portal_block(
        doc,
        title="11.4 Canteen",
        purpose="The Canteen portal supports menu/item information, order creation and review of saved orders.",
        responsibilities=["Maintain menu items where used", "Create and review customer orders", "Track order status", "Review related canteen reports"],
        steps=[
            "Sign in to Canteen Manager.",
            "Open Menu Management if items must be updated.",
            "Open Orders → Add Order.",
            "Enter customer, items, quantity, payment details and save.",
            "Confirm the order appears in the Orders list after saving and after refresh.",
        ],
        info_needed=["Customer type and identity", "Menu items and quantities", "Payment method"],
        approval="Canteen manages its operational orders. Finance processes remain separate unless school policy links them.",
        outcome="Each saved order remains visible in the Orders list for follow-up.",
        mistakes=["Saving without selecting menu items", "Wrong customer identity", "Assuming an order is saved before confirmation appears"],
    )

    portal_block(
        doc,
        title="11.5 IT Support",
        purpose="IT Support records and tracks technology issues, progress and resolutions, and may maintain asset information.",
        responsibilities=["Log support requests", "Update progress", "Record resolution", "Maintain asset notes where available"],
        steps=[
            "Sign in to IT Support Manager.",
            "Open tickets or assets menus.",
            "Create or update a request with clear problem details.",
            "Update status until resolved and record the outcome.",
        ],
        info_needed=["Reporter name", "Device or service affected", "Priority and description"],
        approval="IT owns technical resolution; requesting departments confirm when service is restored.",
        outcome="Tickets show a clear status history.",
        mistakes=["Closing tickets without confirmation", "Missing asset identifiers"],
    )

    portal_block(
        doc,
        title="11.6 Housekeeping",
        purpose="Housekeeping manages cleaning schedules, duty allotments and completion records for campus sanitation.",
        responsibilities=["Maintain schedules", "Assign duties", "Record completion", "Escalate maintenance concerns"],
        steps=[
            "Sign in to Housekeeping Manager.",
            "Open schedules or duty allotments.",
            "Assign tasks and update completion status.",
            "Note concerns for supervisor follow-up.",
        ],
        info_needed=["Area/location", "Shift timing", "Staff assigned"],
        approval="Housekeeping Supervisor/Manager reviews completion; facilities issues may be handed to the relevant department.",
        outcome="Schedules and completion records are available for inspection.",
        mistakes=["Marking incomplete work as done", "Assigning the wrong area"],
    )

    add_heading(doc, "12. Governance and Audit Portals", 1)

    portal_block(
        doc,
        title="12.1 Joint Director",
        purpose="Joint Director provides cross-department oversight, assigned task follow-up and progress review.",
        responsibilities=["Review assigned oversight tasks", "Follow progress across departments", "Request reports and updates"],
        steps=[
            "Sign in to Joint Director.",
            "Open dashboard and assigned task or review menus.",
            "Record follow-up actions and confirm status updates.",
        ],
        info_needed=["Department updates", "Assigned task list"],
        approval="Joint Director oversees; operational approvals remain with department heads unless specifically assigned.",
        outcome="Follow-up items have clear owners and status.",
        mistakes=["Closing follow-up without evidence"],
    )

    portal_block(
        doc,
        title="12.2 Joint Director Audit",
        purpose="Joint Director Audit assigns audits, reviews findings and tracks corrective actions to closure.",
        responsibilities=["Create audit assignments", "Review findings and evidence", "Track corrective actions", "Close audits when criteria are met"],
        steps=[
            "Sign in to Joint Director Audit.",
            "Assign an audit to Process/Quality auditors as required.",
            "Review submitted findings and evidence.",
            "Monitor corrective actions until closure.",
        ],
        info_needed=["Audit scope", "Department under review", "Evidence documents"],
        approval="Auditors raise findings; Joint Director Audit oversees closure discipline.",
        outcome="Assignments, findings and closures are traceable.",
        mistakes=["Closing findings without verified corrective action"],
    )

    portal_block(
        doc,
        title="12.3 Process Auditor",
        purpose="Process Auditors examine whether school processes are followed, record findings and propose corrective actions.",
        responsibilities=["Execute assigned process reviews", "Check evidence", "Register findings", "Follow compliance actions"],
        steps=[
            "Sign in to Process Auditor.",
            "Open assigned fieldwork or audit tasks.",
            "Record evidence-based findings.",
            "Update corrective-action progress.",
        ],
        info_needed=["Assignment brief", "Process checklist", "Evidence"],
        approval="Findings feed Joint Director Audit review.",
        outcome="Process findings are documented for follow-up.",
        mistakes=["Recording opinions without evidence"],
    )

    portal_block(
        doc,
        title="12.4 Quality Auditor",
        purpose="Quality Auditors perform quality checks, record inspection results and non-conformities, and track corrective actions.",
        responsibilities=["Conduct quality inspections", "Log non-conformities", "Track corrective actions", "Support quality review"],
        steps=[
            "Sign in to Quality Auditor.",
            "Open assigned quality checks.",
            "Enter inspection results and non-conformities.",
            "Update corrective actions until verified.",
        ],
        info_needed=["Inspection checklist", "Area or sample under review"],
        approval="Quality findings are reviewed under the audit governance process.",
        outcome="Quality results and corrective actions are recorded.",
        mistakes=["Clearing a non-conformity without re-check"],
    )

    add_heading(doc, "12.5 Note on Hostel-Related Activities", 2)
    p(
        doc,
        "In the current role structure, hostel information for families is viewed through Parent and Student menus "
        "(for example hostel details or hostel-related payments where shown). Front Office maintains hostel gate-pass "
        "records that may include a warden-approval field for Security. There is no separate Hostel Warden login "
        "portal among the standard QMIS ERP roles described in this manual. Day-to-day hostel operational decisions "
        "continue under the school’s administrative arrangements and the Front Office / Security records above.",
    )

    # Workflows
    add_heading(doc, "13. End-to-End School Workflows", 1)

    workflow_block(
        doc,
        title="Workflow 1: Academic Year and School Setup",
        start="Admin (with leadership direction).",
        info="Academic year name and dates; planned classes, sections and capacities.",
        sequence=[
            "Confirm the academic year to be used.",
            "Create or activate classes for that year.",
            "Create sections under each class with capacity.",
            "Share the structure with Principal, Coordinators and Teachers.",
        ],
        review="Leadership confirms the structure before enrolment peaks.",
        next_dept="PRM uses class names during admission; Coordinators/Principal use sections for placement.",
        check="Class Details lists the correct year, classes and sections.",
    )

    workflow_block(
        doc,
        title="Workflow 2: Admission to Student Enrolment",
        start="PRM / Front Office (Admin may assist).",
        info="Child and guardian details, preferred class, contacts, documents as per policy.",
        sequence=[
            "Register the enquiry and follow up.",
            "Create the admission record.",
            "Complete enrolment and guardian information.",
            "Create parent/student portal accounts when required and guide first password change.",
            "Hand the enrolled student to academic placement.",
        ],
        review="Office supervisor checks completeness before handover.",
        next_dept="Principal / Coordinator for section placement.",
        check="Student appears as enrolled; parent/student can sign in if accounts were created.",
    )

    workflow_block(
        doc,
        title="Workflow 3: Student Placement and Teacher Allocation",
        start="Coordinator (grade band) or Principal.",
        info="Enrolled student, target section, capacity, teacher subject allocations.",
        sequence=[
            "Open placement / section assignment.",
            "Assign section (roll number assigned by process).",
            "Principal maintains teacher allocations for classes/subjects.",
            "Teachers see assigned classes in their portal.",
        ],
        review="Principal oversees allocations and placement exceptions.",
        next_dept="Teachers for attendance and teaching; Parents see updated class/section.",
        check="Student shows correct class–section; teacher lists include the class.",
    )

    workflow_block(
        doc,
        title="Workflow 4: Timetable and Lesson Planning",
        start="Principal/Coordinator for timetable; Teacher for lesson plans.",
        info="Periods, subjects, teachers; lesson objectives and dates.",
        sequence=[
            "Prepare class timetable and submit for Director approval where required.",
            "Teachers prepare lesson plans and submit for academic review.",
            "Coordinator → Principal → Director review lesson plans as configured.",
            "Teachers follow the approved timetable.",
        ],
        review="Director for timetable; academic chain for lesson plans.",
        next_dept="Teaching staff deliver lessons; leadership monitors compliance.",
        check="Approved timetable is visible; lesson plans show approved status.",
    )

    workflow_block(
        doc,
        title="Workflow 5: Attendance",
        start="Teacher (class attendance).",
        info="Date, class/section, student list.",
        sequence=[
            "Open attendance for the class.",
            "Mark present/absent (and related statuses).",
            "Save the attendance session.",
            "Authorised staff and parents/students view permitted attendance information.",
        ],
        review="Coordinators/Principal may monitor patterns.",
        next_dept="Families via Parent/Student portals where enabled.",
        check="Attendance for the date is stored and visible to authorised viewers.",
    )

    workflow_block(
        doc,
        title="Workflow 6: Examinations, Marks and Results",
        start="Principal/Coordinator prepare exam timetable; Teachers enter marks.",
        info="Exam name, subjects, dates, max marks, student list.",
        sequence=[
            "Create and submit the examination timetable for Director approval.",
            "After approval, Teachers/Coordinators enter marks in batches.",
            "Submit marks for approval.",
            "Coordinator reviews, then Principal, then Director provides final approval (publication).",
            "Parents and students view published results for the correct learner only.",
        ],
        review="Exam timetable: Director. Marks: Coordinator → Principal → Director.",
        next_dept="Parent and Student portals for published results; notifications may alert parents.",
        check="Only approved/published results appear for families; draft marks remain internal.",
    )

    workflow_block(
        doc,
        title="Workflow 7: Student Promotion and Academic History",
        start="Coordinator / Principal under school promotion policy.",
        info="Current class/section, promotion target, academic standing.",
        sequence=[
            "Review eligible students.",
            "Run promotion / section movement as authorised.",
            "Retain academic history so past results remain attributable to the student.",
        ],
        review="Principal confirms promotion decisions.",
        next_dept="Teachers of the new class; families see updated class information.",
        check="Student’s new class/section is correct; historical results remain accessible.",
    )

    workflow_block(
        doc,
        title="Workflow 8: Student Fees and Payments",
        start="Account Head configures fees; Finance Assistant or Account Head collects.",
        info="Fee structure, student, instalments due, payment mode.",
        sequence=[
            "Account Head sets or updates fee structures.",
            "Staff locate the student ledger and outstanding instalments.",
            "Finance Assistant or Account Head records permitted collection and receipt.",
            "Account Head handles structure changes and finance approvals.",
            "Parents may pay online where enabled and view receipts.",
        ],
        review="Account Head for configuration and approvals; Assistant for routine collection only.",
        next_dept="Parents receive receipts; leadership may review collection reports.",
        check="Payment appears on the student ledger after save; receipt is available.",
    )

    workflow_block(
        doc,
        title="Workflow 9: Parent and Student Access",
        start="PRM/Admin during enrolment when portal accounts are created.",
        info="Unique emails for parent and student; temporary credentials delivered securely by the school.",
        sequence=[
            "Create linked accounts at enrolment.",
            "User signs in and changes password if required.",
            "Parent selects the linked child; student sees only self.",
            "Users open only the menus allowed for their role.",
        ],
        review="Office confirms linkage is correct if a parent reports a missing child.",
        next_dept="Academic and finance data appear automatically once published or posted.",
        check="Parent cannot see unrelated children; student cannot see other learners’ results.",
    )

    workflow_block(
        doc,
        title="Workflow 10: Transport Assignment and Daily Duty",
        start="Transport Manager.",
        info="Routes, vehicles, drivers, student addresses/stop preferences.",
        sequence=[
            "Maintain vehicles, drivers and routes.",
            "Assign students to routes.",
            "Assign driver duties.",
            "Drivers execute duties and report issues.",
            "Parents/students may view transport information where provided.",
        ],
        review="Transport Manager monitors exceptions.",
        next_dept="Drivers for execution; families for visibility.",
        check="Student–route and driver–vehicle assignments match the day’s plan.",
    )

    workflow_block(
        doc,
        title="Workflow 11: Visitor and Gate Management",
        start="Front Office and/or Gate / Security.",
        info="Visitor identity, purpose, host, pass type.",
        sequence=[
            "Create visitor or gate-pass record.",
            "Security verifies and records entry/exit.",
            "Update pass status on exit or completion.",
        ],
        review="Security supervisor / Gate Keeper Manager as applicable.",
        next_dept="Front Office retains administrative pass history.",
        check="Every entry has a matching record and closed status when the visit ends.",
    )

    workflow_block(
        doc,
        title="Workflow 12: Library Management",
        start="Librarian.",
        info="Book details, borrower identity.",
        sequence=["Catalogue books", "Issue to borrower", "Record return", "Review overdue items as per policy"],
        review="Librarian.",
        next_dept="Borrowers (students/staff) receive/return books.",
        check="Issue and return appear on the borrower’s record.",
    )

    workflow_block(
        doc,
        title="Workflow 13: Stores and Inventory",
        start="Stores / Stationery Manager (requests may start with teaching staff).",
        info="Item, quantity, requesting department.",
        sequence=["Receive approved request", "Issue stock", "Record receipts of new stock", "Report discrepancies"],
        review="Store Manager; some requests involve Principal.",
        next_dept="Requesting department receives materials.",
        check="Stock on hand matches receipts minus issues.",
    )

    workflow_block(
        doc,
        title="Workflow 14: Canteen Operations",
        start="Canteen Manager.",
        info="Customer, menu items, quantities, payment details.",
        sequence=[
            "Ensure menu items are available.",
            "Create the order and save.",
            "Confirm the order in the Orders list.",
            "Update order status through preparation and completion as used by the school.",
        ],
        review="Canteen supervisor.",
        next_dept="Customer receives the order; reports summarise sales.",
        check="Saved order remains listed after refresh.",
    )

    workflow_block(
        doc,
        title="Workflow 15: HR and Staff Administration",
        start="HR.",
        info="Employee biodata, role/designation, leave details.",
        sequence=["Create/update employee", "Support onboarding", "Process leave decisions with department heads", "Maintain attendance-related HR records where used"],
        review="HR and department leave approvers.",
        next_dept="Employees and managers see outcomes in their leave/HR views.",
        check="Employee profile and leave status are correct.",
    )

    workflow_block(
        doc,
        title="Workflow 16: IT Support and Housekeeping",
        start="Requester (staff) or IT/Housekeeping operator.",
        info="Location/device, description, priority.",
        sequence=[
            "Log the ticket or housekeeping task.",
            "Assign where the portal supports assignment.",
            "Update progress.",
            "Resolve and record completion.",
            "Requester confirms service restoration when applicable.",
        ],
        review="IT Support Manager / Housekeeping Manager.",
        next_dept="Original requester.",
        check="Status history shows open → in progress → closed.",
    )

    workflow_block(
        doc,
        title="Workflow 17: Process and Quality Audits",
        start="Joint Director Audit (assignment); Process/Quality Auditors (fieldwork).",
        info="Scope, checklist, evidence.",
        sequence=[
            "Assign the audit.",
            "Auditors collect evidence and register findings.",
            "Agree corrective actions.",
            "Verify and close.",
        ],
        review="Joint Director Audit.",
        next_dept="Audited department implements corrections.",
        check="Every finding has evidence and a closure note.",
    )

    add_heading(doc, "14. How Departments Connect", 1)
    p(
        doc,
        "School operations in QMIS ERP are connected. The following business flow summarises the main handovers:",
    )
    bullets(
        doc,
        [
            "Admission (PRM/Admin) creates enrolled students and may link parent/student accounts.",
            "Academic placement (Coordinator/Principal) assigns class and section.",
            "Teacher allocation and timetable guide classroom teaching and attendance.",
            "Examinations and marks approvals lead to published results.",
            "Published academic information appears for the correct Parent and Student viewers.",
            "Finance maintains fee structures and collections for students/families.",
            "Transport assignments connect students with routes, vehicles and drivers.",
            "HR maintains staff information used across authorised staff processes.",
            "Campus operations (Gate, Library, Stores, Canteen, IT, Housekeeping) keep their service records and hand issues to the responsible person.",
            "Audit roles review assigned areas, record findings and track corrective action.",
        ],
    )
    p(doc, "Simple operational flow:", bold=True)
    p(
        doc,
        "Enquiry → Admission/Enrolment → Placement → Teaching & Attendance → Examinations → "
        "Marks Approval → Published Results → Parent/Student View",
    )
    p(
        doc,
        "Alongside academics: Fees (Finance) · Transport · Gate/Library/Stores/Canteen/IT/HK · HR · Audit oversight",
    )

    add_heading(doc, "15. General Operating Instructions and Troubleshooting", 1)
    add_heading(doc, "15.1 Daily good practice", 2)
    bullets(
        doc,
        [
            "Sign in with your own account and sign out when finished.",
            "Check academic year, class, section, student name and amounts before saving.",
            "After important saves, refresh or reopen the list to confirm the record.",
            "Use remarks when returning work for correction so the next person knows what to fix.",
            "Protect personal data of students, parents and staff.",
        ],
    )
    add_heading(doc, "15.2 If something seems unavailable", 2)
    bullets(
        doc,
        [
            "Confirm you are signed in with the correct role.",
            "Ask the owning department—access is intentionally limited by responsibility.",
            "For password problems, contact the school administrator who manages accounts.",
            "For incorrect academic or fee data, contact the department that created the record.",
        ],
    )
    add_heading(doc, "15.3 Approvals", 2)
    p(
        doc,
        "Many academic items move through a defined review chain. If your item is pending, wait for the "
        "next role or respond promptly when you receive a return remark. Do not create duplicate submissions "
        "unless the school process requires a new submission after denial.",
    )

    add_heading(doc, "16. Closing Note", 1)
    p(
        doc,
        "This Overall User Manual describes how QMIS ERP supports Queen Mira International School’s "
        "administrative, academic, finance, transport, campus operations and audit work. Follow your "
        "school’s local policies for documents, fees, transport safety and child protection in addition "
        "to the steps in this guide.",
    )
    p(doc, "For account access or process clarification, contact your school’s ERP administrator or the department head responsible for the process.", bold=False)

    doc.save(DOCX_PATH)
    print("Wrote", DOCX_PATH)
    return DOCX_PATH


def try_pdf(docx_path: Path):
    # Attempt LibreOffice, then docx2pdf
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT_DIR), str(docx_path)],
            check=False,
            capture_output=True,
        )
        if PDF_PATH.exists():
            print("Wrote", PDF_PATH)
            return PDF_PATH
    try:
        from docx2pdf import convert

        convert(str(docx_path), str(PDF_PATH))
        if PDF_PATH.exists():
            print("Wrote", PDF_PATH)
            return PDF_PATH
    except Exception as exc:  # noqa: BLE001
        print("PDF not generated:", exc)
    return None


if __name__ == "__main__":
    path = build()
    try_pdf(path)
